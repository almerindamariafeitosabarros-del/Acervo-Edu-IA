"""Extração do texto dos documentos para o assistente de IA."""

import os


class TextExtractionError(Exception):
    """O texto do documento não pôde ser lido."""


NO_TEXT_MESSAGE = 'Não foi possível ler o texto deste documento.'
UNSUPPORTED_MESSAGE = 'Formato não suportado pelo assistente. Use PDF, DOCX ou TXT.'


def extract_pdf(file_obj):
    from pypdf import PdfReader

    try:
        reader = PdfReader(file_obj)
        pages = [page.extract_text() or '' for page in reader.pages]
    except Exception as exc:  # arquivo corrompido ou protegido
        raise TextExtractionError(NO_TEXT_MESSAGE) from exc
    return '\n'.join(pages)


def extract_docx(file_obj):
    from docx import Document as DocxDocument

    try:
        document = DocxDocument(file_obj)
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(' | '.join(cell.text for cell in row.cells))
    except Exception as exc:
        raise TextExtractionError(NO_TEXT_MESSAGE) from exc
    return '\n'.join(parts)


def extract_txt(file_obj):
    raw = file_obj.read()
    if isinstance(raw, bytes):
        for encoding in ('utf-8', 'latin-1'):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise TextExtractionError(NO_TEXT_MESSAGE)
    return raw


EXTRACTORS = {
    'pdf': extract_pdf,
    'docx': extract_docx,
    'txt': extract_txt,
}


def extract_text(file_obj, filename):
    """Devolve o texto do arquivo ou levanta TextExtractionError."""
    extension = os.path.splitext(filename)[1].lower().lstrip('.')
    extractor = EXTRACTORS.get(extension)
    if extractor is None:
        raise TextExtractionError(UNSUPPORTED_MESSAGE)
    text = (extractor(file_obj) or '').strip()
    if not text:
        raise TextExtractionError(NO_TEXT_MESSAGE)
    return text
